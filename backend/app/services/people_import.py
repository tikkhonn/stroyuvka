from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import Person, Unit
from app.schemas import (
    RosterImportPreview,
    RosterImportResult,
    RosterParseError,
    RosterParseRow,
)
from app.services.people import (
    FIO_REQUIRED_MSG,
    create_person,
    display_last_name,
    find_match,
    format_display_name,
    format_rank,
    fio_key,
    list_people,
    normalize_department_code,
    parse_fio,
    sync_unit_strength_from_people,
    validate_rank,
)

RANK_HEADERS = {"звание", "воинское звание", "зв", "зв.", "rank"}
FIO_HEADERS = {
    "фио",
    "fio",
    "full_name",
    "фамилия и.о.",
    "фамилия и. о.",
    "фамилия и.о",
    "фамилия и о",
}
LAST_HEADERS = {"фамилия", "фамилии", "last_name", "lastname", "фамилия имя"}
INIT_HEADERS = {"инициалы", "инициал", "и.", "initials", "имя"}
DEPT_HEADERS = {"кафедра", "каф", "каф.", "department", "dept"}

KNOWN_RANKS = sorted(
    {
        "рядовой",
        "ефрейтор",
        "младший сержант",
        "сержант",
        "старший сержант",
        "старшина",
        "прапорщик",
        "старший прапорщик",
        "младший лейтенант",
        "лейтенант",
        "старший лейтенант",
        "капитан",
        "майор",
        "подполковник",
        "полковник",
        "генерал-майор",
        "курсант",
        "к-т",
        "к-н",
        "ст. лейтенант",
        "ст.л-т",
        "мл. лейтенант",
        "мл.л-т",
    },
    key=len,
    reverse=True,
)


@dataclass
class ParsedLine:
    row_number: int
    rank: str
    last_name: str
    first_name: str
    middle_name: str
    source: str
    department_code: str | None = None
    warnings: list[str] = field(default_factory=list)

    @property
    def full_name(self) -> str:
        return format_display_name(self.last_name, self.first_name, self.middle_name)


def _norm_header(value: object) -> str:
    text = str(value or "").strip().casefold()
    return re.sub(r"\s+", " ", text).rstrip(":")


def _is_number_header(name: str) -> bool:
    if not name:
        return False
    if name in {"№", "n", "nn", "no", "num", "number"}:
        return True
    return "п/п" in name or name.startswith("№")


def _is_rank_header(name: str) -> bool:
    return name in RANK_HEADERS or "зван" in name


def _is_fio_header(name: str) -> bool:
    if name in FIO_HEADERS:
        return True
    if "фио" in name or "отчество" in name:
        return True
    return "фамилия" in name and "инициал" not in name


def _is_init_header(name: str) -> bool:
    return name in INIT_HEADERS or "инициал" in name


def _is_dept_header(name: str) -> bool:
    if name in DEPT_HEADERS:
        return True
    return "кафедр" in name


def _looks_like_department(value: str) -> bool:
    text = value.strip()
    if not text:
        return False
    cleaned = re.sub(r"[^\w]", "", text, flags=re.UNICODE)
    return bool(cleaned) and len(cleaned) <= 10


def _looks_like_row_number(value: str) -> bool:
    text = value.strip()
    if not text:
        return False
    cleaned = text.lstrip("№").strip().rstrip(".")
    return cleaned.isdigit()


def _cell_looks_like_header(cell: str) -> bool:
    name = _norm_header(cell)
    if not name:
        return False
    return (
        _is_number_header(name)
        or _is_rank_header(name)
        or _is_fio_header(name)
        or _is_init_header(name)
        or _is_dept_header(name)
        or name in LAST_HEADERS
    )


def _is_header_row(row: list[str]) -> bool:
    if not any(cell.strip() for cell in row):
        return False
    if _map_headers(row) is not None:
        return True
    return any(_cell_looks_like_header(cell) for cell in row if cell.strip())


def _cells_from_excel(content: bytes) -> list[list[str]]:
    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    try:
        ws = wb.active
        rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            rows.append(["" if c is None else str(c).strip() for c in row])
        return rows
    finally:
        wb.close()


def _cells_from_docx(content: bytes) -> tuple[list[list[str]], list[str]]:
    from docx import Document

    doc = Document(io.BytesIO(content))
    tables: list[list[str]] = []
    for table in doc.tables:
        for row in table.rows:
            tables.append([cell.text.strip() for cell in row.cells])
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    return tables, paragraphs


def _map_headers(row: list[str]) -> dict[str, int] | None:
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(row):
        name = _norm_header(cell)
        if _is_number_header(name):
            continue
        if _is_rank_header(name) and "rank" not in mapping:
            mapping["rank"] = idx
        elif _is_init_header(name):
            mapping["init"] = idx
        elif _is_fio_header(name) and "fio" not in mapping:
            mapping["fio"] = idx
        elif (name in LAST_HEADERS or name.startswith("фамилия")) and "last" not in mapping:
            mapping["last"] = idx
        elif _is_dept_header(name) and "dept" not in mapping:
            mapping["dept"] = idx

    if {"rank", "dept", "fio"} <= set(mapping):
        return mapping
    if {"rank", "fio"} <= set(mapping):
        return mapping
    if {"rank", "last", "init"} <= set(mapping):
        return mapping
    if {"rank", "last"} <= set(mapping) and "init" not in mapping:
        mapping["fio"] = mapping.pop("last")
        return mapping
    return None


def _default_mapping(row: list[str]) -> dict[str, int]:
    non_empty = sum(1 for cell in row if cell.strip())
    if non_empty >= 4 and _looks_like_row_number(row[0] if row else ""):
        if len(row) >= 4 and _looks_like_department(row[2] if len(row) > 2 else ""):
            return {"rank": 1, "dept": 2, "fio": 3}
        return {"rank": 1, "fio": 2}
    if non_empty >= 3 and _looks_like_department(row[1] if len(row) > 1 else ""):
        return {"rank": 0, "dept": 1, "fio": 2}
    if non_empty >= 3:
        return {"rank": 0, "last": 1, "init": 2}
    return {"rank": 0, "fio": 1}


def _parse_department_cell(mapping: dict[str, int], row: list[str]) -> str | None:
    if "dept" not in mapping:
        return None
    idx = mapping["dept"]
    raw = (row[idx] if idx < len(row) else "").strip()
    return normalize_department_code(raw)


def _parse_name_cells(
    mapping: dict[str, int], row: list[str]
) -> tuple[str, str, str, list[str]]:
    warnings: list[str] = []
    if "fio" in mapping:
        raw = (row[mapping["fio"]] if mapping["fio"] < len(row) else "").strip()
        if not raw:
            return "", "", "", warnings
        try:
            last_name, first_name, middle_name = parse_fio(raw)
            return last_name, first_name, middle_name, warnings
        except ValueError as exc:
            warnings.append(str(exc))
            return "", "", "", warnings

    last = (row[mapping["last"]] if mapping["last"] < len(row) else "").strip()
    init = (row[mapping["init"]] if mapping["init"] < len(row) else "").strip()
    if not last:
        return "", "", "", warnings
    combined = f"{last} {init}".strip() if init else last
    try:
        last_name, first_name, middle_name = parse_fio(combined)
        return last_name, first_name, middle_name, warnings
    except ValueError as exc:
        warnings.append(str(exc))
        return "", "", "", warnings


def _parse_table_rows(rows: list[list[str]], source: str) -> list[ParsedLine]:
    if not rows:
        return []
    start = 0
    mapping = _map_headers(rows[0])
    if mapping:
        start = 1
    else:
        mapping = _default_mapping(rows[0])

    parsed: list[ParsedLine] = []
    for offset, row in enumerate(rows[start:], start=start + 1):
        if not any(row):
            continue
        if _is_header_row(row):
            continue
        raw_rank = (row[mapping["rank"]] if mapping["rank"] < len(row) else "").strip()
        rank = ""
        rank_warnings: list[str] = []
        if raw_rank:
            try:
                rank = validate_rank(raw_rank)
            except ValueError as exc:
                rank_warnings.append(str(exc))
        department_code = _parse_department_cell(mapping, row)
        last_name, first_name, middle_name, name_warnings = _parse_name_cells(mapping, row)
        if not last_name and not rank and not first_name:
            continue
        warnings = list(name_warnings) + rank_warnings
        if not raw_rank:
            warnings.append("Нет звания")
        if name_warnings and not last_name:
            parsed.append(
                ParsedLine(
                    row_number=offset,
                    rank=rank,
                    last_name="",
                    first_name="",
                    middle_name="",
                    source=source,
                    department_code=department_code,
                    warnings=warnings,
                )
            )
            continue
        parsed.append(
            ParsedLine(
                row_number=offset,
                rank=rank,
                last_name=last_name,
                first_name=first_name,
                middle_name=middle_name,
                source=source,
                department_code=department_code,
                warnings=warnings,
            )
        )
    return parsed


def _parse_loose_line(text: str, row_number: int) -> ParsedLine | None:
    line = re.sub(r"\s+", " ", text).strip()
    if not line:
        return None
    if _is_header_row([line]):
        return None
    skip_headers = RANK_HEADERS | LAST_HEADERS | INIT_HEADERS | FIO_HEADERS
    if _norm_header(line) in skip_headers:
        return None
    if _is_number_header(_norm_header(line)):
        return None
    if _is_rank_header(_norm_header(line)) or _is_fio_header(_norm_header(line)):
        return None
    if _norm_header(line) in {
        "звание фамилия инициалы",
        "звание фамилия и.о.",
        "воинское звание фамилия инициалы",
        "воинское звание фамилия и.о.",
    }:
        return None

    lower = line.casefold()
    rank = ""
    rest = line
    for known in KNOWN_RANKS:
        if lower.startswith(known):
            rank = line[: len(known)].strip()
            rest = line[len(known) :].strip()
            break

    if rank and rest:
        try:
            validated_rank = validate_rank(rank)
            last_name, first_name, middle_name = parse_fio(rest)
            return ParsedLine(
                row_number=row_number,
                rank=validated_rank,
                last_name=last_name,
                first_name=first_name,
                middle_name=middle_name,
                source="text",
            )
        except ValueError as exc:
            return ParsedLine(
                row_number=row_number,
                rank="",
                last_name="",
                first_name="",
                middle_name="",
                source="text",
                warnings=[str(exc)],
            )

    parts = line.split(None, 1)
    if len(parts) == 2:
        try:
            validated_rank = validate_rank(parts[0].strip())
            last_name, first_name, middle_name = parse_fio(parts[1])
            return ParsedLine(
                row_number=row_number,
                rank=validated_rank,
                last_name=last_name,
                first_name=first_name,
                middle_name=middle_name,
                source="text",
            )
        except ValueError as exc:
            return ParsedLine(
                row_number=row_number,
                rank="",
                last_name="",
                first_name="",
                middle_name="",
                source="text",
                warnings=[str(exc)],
            )

    return ParsedLine(
        row_number=row_number,
        rank="",
        last_name="",
        first_name="",
        middle_name="",
        source="text",
        warnings=[f"Не удалось разобрать строку: {line}"],
    )


def parse_roster_file(filename: str, content: bytes) -> tuple[list[ParsedLine], list[RosterParseError]]:
    name = (filename or "").lower()
    errors: list[RosterParseError] = []
    parsed: list[ParsedLine] = []
    try:
        if name.endswith(".xlsx") or name.endswith(".xlsm"):
            parsed = _parse_table_rows(_cells_from_excel(content), "xlsx")
        elif name.endswith(".docx"):
            table_rows, paragraphs = _cells_from_docx(content)
            if table_rows:
                parsed = _parse_table_rows(table_rows, "docx-table")
            if not parsed:
                for i, paragraph in enumerate(paragraphs, start=1):
                    line = _parse_loose_line(paragraph, i)
                    if line:
                        parsed.append(line)
        else:
            errors.append(
                RosterParseError(message="Нужен файл Excel (.xlsx) или Word (.docx)")
            )
            return [], errors
    except Exception as exc:  # noqa: BLE001 — разбор пользовательского файла
        errors.append(RosterParseError(message=f"Не удалось прочитать файл: {exc}"))
        return [], errors

    usable: list[ParsedLine] = []
    for item in parsed:
        if item.warnings:
            errors.append(
                RosterParseError(row_number=item.row_number, message=item.warnings[0])
            )
            continue
        if not item.last_name or not item.first_name or not item.middle_name:
            errors.append(
                RosterParseError(
                    row_number=item.row_number,
                    message=FIO_REQUIRED_MSG,
                )
            )
            continue
        if not item.rank:
            errors.append(
                RosterParseError(row_number=item.row_number, message="Нет звания")
            )
            continue
        usable.append(item)

    seen: dict[tuple[str, str, str], int] = {}
    for item in usable:
        key = fio_key(item.last_name, item.first_name, item.middle_name)
        if key in seen:
            errors.append(
                RosterParseError(
                    row_number=item.row_number,
                    message=(
                        f"В файле повтор: {item.full_name} "
                        f"(строка {seen[key]})"
                    ),
                )
            )
        else:
            seen[key] = item.row_number
    return usable, errors


async def preview_import(
    session: AsyncSession,
    unit: Unit,
    filename: str,
    content: bytes,
    mode: str,
) -> RosterImportPreview:
    if mode not in ("replace", "upsert"):
        raise ValueError("mode должен быть replace или upsert")
    parsed, errors = parse_roster_file(filename, content)
    existing = await list_people(session, unit.id, include_inactive=True)
    rows: list[RosterParseRow] = []
    to_add = to_update = to_restore = 0
    matched_ids: set[int] = set()

    for item in parsed:
        active, inactive = find_match(
            existing, item.last_name, item.first_name, item.middle_name
        )
        action = "add"
        person_id = None
        warnings = list(item.warnings)
        if len(active) > 1 or (not active and len(inactive) > 1):
            errors.append(
                RosterParseError(
                    row_number=item.row_number,
                    message=f"Конфликт: несколько записей {item.full_name}",
                )
            )
            action = "conflict"
        elif len(active) == 1:
            action = "update"
            person_id = active[0].id
            matched_ids.add(active[0].id)
            to_update += 1
        elif len(inactive) == 1:
            action = "restore"
            person_id = inactive[0].id
            matched_ids.add(inactive[0].id)
            to_restore += 1
        else:
            to_add += 1
        rows.append(
            RosterParseRow(
                row_number=item.row_number,
                rank=item.rank,
                full_name=item.full_name,
                last_name=item.last_name,
                first_name=item.first_name,
                middle_name=item.middle_name,
                department_code=item.department_code,
                source=item.source,
                action=action,
                person_id=person_id,
                warnings=warnings,
            )
        )

    to_deactivate = 0
    if mode == "replace":
        for person in existing:
            if person.is_active and person.id not in matched_ids:
                to_deactivate += 1
                rows.append(
                    RosterParseRow(
                        row_number=0,
                        rank=format_rank(person.rank),
                        full_name=display_last_name(person),
                        last_name=person.last_name,
                        first_name=person.first_name,
                        middle_name=person.middle_name,
                        department_code=person.department_code,
                        source="db",
                        action="deactivate",
                        person_id=person.id,
                    )
                )

    return RosterImportPreview(
        rows=rows,
        errors=errors,
        to_add=to_add,
        to_update=to_update,
        to_restore=to_restore,
        to_deactivate=to_deactivate,
        can_apply=len(errors) == 0 and len(parsed) > 0,
    )


async def apply_import(
    session: AsyncSession,
    unit: Unit,
    filename: str,
    content: bytes,
    mode: str,
) -> RosterImportResult:
    preview = await preview_import(session, unit, filename, content, mode)
    if not preview.can_apply:
        if preview.errors:
            raise ValueError(preview.errors[0].message)
        raise ValueError("В файле нет строк для импорта")

    existing = await list_people(session, unit.id, include_inactive=True)
    added = updated = restored = deactivated = 0
    matched_ids: set[int] = set()

    for row in preview.rows:
        if row.action == "deactivate":
            continue
        active, inactive = find_match(
            existing, row.last_name, row.first_name, row.middle_name
        )
        if row.action == "add":
            person = await create_person(
                session,
                unit,
                row.rank,
                row.full_name,
                middle_name=row.middle_name,
                department_code=row.department_code,
            )
            existing.append(person)
            matched_ids.add(person.id)
            added += 1
        elif row.action == "update" and active:
            person = active[0]
            person.rank = row.rank
            person.last_name = row.last_name
            person.first_name = row.first_name
            person.middle_name = row.middle_name or None
            person.department_code = normalize_department_code(row.department_code)
            person.is_active = True
            matched_ids.add(person.id)
            updated += 1
        elif row.action == "restore" and inactive:
            person = inactive[0]
            person.rank = row.rank
            person.last_name = row.last_name
            person.first_name = row.first_name
            person.middle_name = row.middle_name or None
            person.department_code = normalize_department_code(row.department_code)
            person.is_active = True
            matched_ids.add(person.id)
            restored += 1

    if mode == "replace":
        for person in existing:
            if person.is_active and person.id not in matched_ids:
                person.is_active = False
                deactivated += 1

    await session.flush()
    total = await sync_unit_strength_from_people(session, unit.id)
    return RosterImportResult(
        added=added,
        updated=updated,
        restored=restored,
        deactivated=deactivated,
        total_list=total,
    )
