"""Excel для ДПА: общий расход, переменный состав, список больных."""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.enums import AbsenceCategoryCode, UnitType
from app.models import AbsenceEntry, Hospital, Unit
from app.schemas import AttendanceAggregate, ChessboardSickEntry
from app.services.attendance import (
    aggregate_from_entries,
    compute_academy_aggregate,
    empty_aggregate,
    get_unit_strength,
    list_absence_entries,
    sum_aggregates,
    _normalize_code,
)
from app.services.org import get_courses_for_faculty, get_courses_for_location
from app.services.reports import build_chessboard_sick_summary
from app.services.unit_ids import LOCATION_LEKHTUSI, LOCATION_PUSHKIN

SICK_HEADERS = ("№", "Подр", "Фамилия, инициалы", "Мед. учреждение", "Диагноз", "Дата")
SICK_COL_WIDTHS = (6, 10, 28, 24, 28, 14)

ACADEMY_HEADERS = (
    "Подразделение",
    "П/С",
    "Н/Л",
    "Н",
    "К",
    "О",
    "Б",
    "Арест",
    "Караул",
    "СОЧ",
)


@dataclass
class _CourseSnap:
    agg: AttendanceAggregate
    sick_by: dict[int | None, int]


@dataclass
class _RashodLine:
    name: str
    agg: AttendanceAggregate
    sick_by: dict[int | None, int] = field(default_factory=dict)
    bold: bool = False


def _date_cell(value: date) -> str:
    return f"{value.day:02d}.{value.month:02d}.{value.year}"


def _num(value: int) -> int | None:
    return value if value else None


def _thin_border() -> Border:
    side = Side(style="thin", color="000000")
    return Border(left=side, right=side, top=side, bottom=side)


def _styles() -> tuple[Border, Font, Font, Alignment, Alignment]:
    thin = _thin_border()
    return (
        thin,
        Font(bold=True, size=12),
        Font(bold=True),
        Alignment(horizontal="center", vertical="center", wrap_text=True),
        Alignment(horizontal="left", vertical="center", wrap_text=True),
    )


def _is_officer(entry: ChessboardSickEntry) -> bool:
    return (entry.unit_name or "").startswith("Офицеры") or entry.location_name == "Офицеры"


def _podr(entry: ChessboardSickEntry) -> str:
    if _is_officer(entry):
        faculty_id = entry.faculty_id or entry.unit_id
        return f"оф.{faculty_id}"
    return str(entry.unit_id)


def _facility_name(entry: ChessboardSickEntry) -> str:
    name = (entry.hospital_name or "").strip()
    if not name or name == "Не указана":
        return "Не указано"
    return name


def _sick_sort_key(entry: ChessboardSickEntry) -> tuple:
    if _is_officer(entry):
        return (1, entry.faculty_id or 0, (entry.last_name or "").lower(), entry.id)
    return (0, entry.unit_id, (entry.last_name or "").lower(), entry.id)


def _merge_sick(parts: Iterable[dict[int | None, int]]) -> dict[int | None, int]:
    merged: dict[int | None, int] = {}
    for part in parts:
        for key, count in part.items():
            merged[key] = merged.get(key, 0) + count
    return merged


def _sick_by_hospital(entries: list[AbsenceEntry]) -> dict[int | None, int]:
    counts: dict[int | None, int] = {}
    for entry in entries:
        if _normalize_code(entry.category_code) != AbsenceCategoryCode.SICK:
            continue
        key = entry.hospital_id
        counts[key] = counts.get(key, 0) + 1
    return counts


def _write_title(ws: Worksheet, title: str, last_col: int, row: int = 1) -> None:
    thin, title_font, _header_font, center, _left = _styles()
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    cell = ws.cell(row, 1, title)
    cell.font = title_font
    cell.alignment = center
    ws.row_dimensions[row].height = 28
    for col in range(1, last_col + 1):
        ws.cell(row, col).border = thin


def _write_header_row(ws: Worksheet, headers: tuple[str, ...] | list[str], row: int) -> None:
    thin, _title_font, header_font, center, _left = _styles()
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row, col, header)
        cell.font = header_font
        cell.alignment = center
        cell.border = thin


def _set_widths(ws: Worksheet, widths: Iterable[float]) -> None:
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width


def _write_academy_sheet(ws: Worksheet, agg: AttendanceAggregate, report_date: date) -> None:
    ws.title = "Общий расход"
    headers = ACADEMY_HEADERS
    _write_title(
        ws,
        f"Расход личного состава на {_date_cell(report_date)}",
        len(headers),
    )
    _write_header_row(ws, headers, 2)
    present_for_report = agg.present + agg.dismissal + agg.away_dorm + agg.other
    values = [
        "Всего за академию",
        _num(agg.total_list),
        _num(present_for_report),
        _num(agg.duty),
        _num(agg.trip),
        _num(agg.leave),
        _num(agg.sick),
        _num(agg.arrest),
        None,
        None,
    ]
    thin, _title_font, header_font, center, left = _styles()
    for col, value in enumerate(values, start=1):
        cell = ws.cell(3, col, value)
        cell.border = thin
        cell.alignment = left if col == 1 else center
        if col == 1:
            cell.font = header_font
    _set_widths(ws, (22, 8, 8, 8, 8, 8, 8, 10, 10, 8))
    ws.freeze_panes = "A3"


def _variable_headers(facilities: list[tuple[int | None, str]]) -> list[str]:
    return [
        "Подразделение",
        "По списку",
        "На лицо",
        "Наряд",
        *[name for _hid, name in facilities],
        "Отпуск",
        "Командировка",
        "Вне общежития",
        "Увольнение",
        "Караул",
        "Арест",
        "Проверка",
    ]


def _line_values(line: _RashodLine, facilities: list[tuple[int | None, str]]) -> list[object]:
    sick_vals = [line.sick_by.get(hid, 0) for hid, _name in facilities]
    mapped_sick = sum(sick_vals)
    leftover_sick = max(0, line.agg.sick - mapped_sick)
    check = (
        line.agg.present
        + line.agg.duty
        + line.agg.trip
        + line.agg.leave
        + mapped_sick
        + leftover_sick
        + line.agg.dismissal
        + line.agg.away_dorm
        + line.agg.other
        + line.agg.arrest
    )
    raw = [
        line.name,
        line.agg.total_list,
        line.agg.present,
        line.agg.duty,
        *sick_vals,
        line.agg.leave,
        line.agg.trip,
        line.agg.away_dorm + line.agg.other,
        line.agg.dismissal,
        0,
        line.agg.arrest,
        check,
    ]
    out: list[object] = []
    for idx, value in enumerate(raw):
        if idx == 0:
            out.append(value)
        else:
            out.append(_num(int(value)))
    return out


def _write_data_row(
    ws: Worksheet, row: int, values: list[object], *, bold: bool = False
) -> None:
    thin, _title_font, header_font, center, left = _styles()
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row, col, value)
        cell.border = thin
        cell.alignment = left if col == 1 else center
        if bold:
            cell.font = header_font


def _sum_lines(name: str, lines: list[_RashodLine]) -> _RashodLine:
    if not lines:
        return _RashodLine(name=name, agg=empty_aggregate(), sick_by={}, bold=True)
    return _RashodLine(
        name=name,
        agg=sum_aggregates([line.agg for line in lines]),
        sick_by=_merge_sick(line.sick_by for line in lines),
        bold=True,
    )


def _write_variable_sheet(
    ws: Worksheet,
    report_date: date,
    facilities: list[tuple[int | None, str]],
    faculty_lines: list[_RashodLine],
    lekhtusi_lines: list[_RashodLine],
    pushkin_lines: list[_RashodLine],
) -> None:
    ws.title = "Переменный состав"
    headers = _variable_headers(facilities)
    last_col = len(headers)
    _write_title(
        ws,
        "Расход переменного состава ВКА имени А.Ф. Можайского по состоянию на "
        f"{_date_cell(report_date)}",
        last_col,
    )

    row = 2
    _write_header_row(ws, headers, row)
    row += 1
    for line in faculty_lines:
        _write_data_row(ws, row, _line_values(line, facilities), bold=line.bold)
        row += 1
    _write_data_row(
        ws, row, _line_values(_sum_lines("Итого", faculty_lines), facilities), bold=True
    )
    row += 2

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    ws.cell(row, 1, "Лехтуси")
    ws.cell(row, 1).font = Font(bold=True)
    row += 1
    _write_header_row(ws, headers, row)
    row += 1
    for line in lekhtusi_lines:
        _write_data_row(ws, row, _line_values(line, facilities), bold=line.bold)
        row += 1
    _write_data_row(
        ws,
        row,
        _line_values(_sum_lines("Итого за Лехтуси", lekhtusi_lines), facilities),
        bold=True,
    )
    row += 2

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    ws.cell(row, 1, "в том числе Пушкин")
    ws.cell(row, 1).font = Font(bold=True)
    row += 1
    _write_header_row(ws, headers, row)
    row += 1
    for line in pushkin_lines:
        _write_data_row(ws, row, _line_values(line, facilities), bold=line.bold)
        row += 1
    _write_data_row(
        ws,
        row,
        _line_values(_sum_lines("Итого за Пушкин", pushkin_lines), facilities),
        bold=True,
    )

    widths = [22, 12, 10, 10]
    widths.extend(14 for _ in facilities)
    widths.extend([12, 14, 14, 12, 10, 10, 12])
    _set_widths(ws, widths)
    ws.freeze_panes = "A3"


def _write_sick_sheet(ws: Worksheet, report_date: date, entries: list[ChessboardSickEntry]) -> None:
    ws.title = "Больные все"
    thin, _title_font, _header_font, center, left = _styles()
    _write_title(
        ws,
        "Список военнослужащих ВКА, находящихся на лечении на "
        f"{_date_cell(report_date)}",
        6,
    )
    _write_header_row(ws, SICK_HEADERS, 2)

    rows = sorted(entries, key=_sick_sort_key)
    for index, entry in enumerate(rows, start=1):
        row_idx = index + 2
        values = (
            index,
            _podr(entry),
            entry.last_name,
            _facility_name(entry),
            (entry.note or "").strip(),
            _date_cell(entry.status_date),
        )
        aligns = (center, center, left, left, left, center)
        for col, (value, alignment) in enumerate(zip(values, aligns), start=1):
            cell = ws.cell(row_idx, col, value)
            cell.alignment = alignment
            cell.border = thin

    _set_widths(ws, SICK_COL_WIDTHS)
    ws.freeze_panes = "A3"
    ws.print_title_rows = "1:2"


def _line_from_snap(name: str, snap: _CourseSnap) -> _RashodLine:
    return _RashodLine(name=name, agg=snap.agg, sick_by=snap.sick_by)


async def _load_course_snaps(
    session: AsyncSession, courses: list[Unit], report_date: date
) -> dict[int, _CourseSnap]:
    snaps: dict[int, _CourseSnap] = {}
    for course in courses:
        if course.id in snaps:
            continue
        entries = await list_absence_entries(session, course.id, report_date)
        total = await get_unit_strength(session, course.id)
        snaps[course.id] = _CourseSnap(
            agg=aggregate_from_entries(total, entries),
            sick_by=_sick_by_hospital(entries),
        )
    return snaps


def _facility_columns(
    hospitals: list[Hospital], used_ids: set[int | None]
) -> list[tuple[int | None, str]]:
    cols: list[tuple[int | None, str]] = []
    for hospital in hospitals:
        if hospital.is_active or hospital.id in used_ids:
            cols.append((hospital.id, hospital.name))
    if None in used_ids:
        cols.append((None, "Не указано"))
    return cols


async def build_sick_xlsx(session: AsyncSession, report_date: date) -> bytes:
    academy = await compute_academy_aggregate(session, report_date)
    sick_summary = await build_chessboard_sick_summary(session, report_date)

    faculties = list(
        (
            await session.execute(
                select(Unit)
                .where(Unit.type == UnitType.FACULTY, Unit.is_active.is_(True))
                .order_by(Unit.id)
            )
        )
        .scalars()
        .all()
    )
    all_courses: list[Unit] = []
    faculty_courses: dict[int, list[Unit]] = {}
    for faculty in faculties:
        courses = await get_courses_for_faculty(session, faculty.id)
        faculty_courses[faculty.id] = courses
        all_courses.extend(courses)

    seen = {c.id for c in all_courses}
    lekhtusi_courses = await get_courses_for_location(session, LOCATION_LEKHTUSI)
    pushkin_courses = await get_courses_for_location(session, LOCATION_PUSHKIN)
    extra = [c for c in lekhtusi_courses + pushkin_courses if c.id not in seen]
    snaps = await _load_course_snaps(session, all_courses + extra, report_date)

    used_ids: set[int | None] = set()
    for snap in snaps.values():
        used_ids.update(snap.sick_by.keys())

    hospitals = list(
        (
            await session.execute(
                select(Hospital).order_by(Hospital.sort_order, Hospital.name, Hospital.id)
            )
        )
        .scalars()
        .all()
    )
    facilities = _facility_columns(hospitals, used_ids)

    faculty_lines: list[_RashodLine] = []
    for faculty in faculties:
        parts = [
            snaps[c.id] for c in faculty_courses[faculty.id] if c.id in snaps
        ]
        if parts:
            faculty_lines.append(
                _RashodLine(
                    name=faculty.name,
                    agg=sum_aggregates([p.agg for p in parts]),
                    sick_by=_merge_sick(p.sick_by for p in parts),
                )
            )
        else:
            faculty_lines.append(
                _RashodLine(name=faculty.name, agg=empty_aggregate(), sick_by={})
            )

    lekhtusi_lines = [
        _line_from_snap(str(c.id), snaps[c.id])
        for c in lekhtusi_courses
        if c.id in snaps
    ]
    pushkin_lines = [
        _line_from_snap(str(c.id), snaps[c.id])
        for c in pushkin_courses
        if c.id in snaps
    ]

    wb = Workbook()
    _write_academy_sheet(wb.active, academy, report_date)
    _write_variable_sheet(
        wb.create_sheet(),
        report_date,
        facilities,
        faculty_lines,
        lekhtusi_lines,
        pushkin_lines,
    )
    _write_sick_sheet(wb.create_sheet(), report_date, sick_summary.entries)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
